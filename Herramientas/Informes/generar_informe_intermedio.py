"""Generar DOCX y HTML autónomo desde el informe Markdown y sus imágenes locales.

Requiere python-docx y Pillow. No consulta CVAT, no ejecuta modelos y no usa test.
El PDF se exporta desde el DOCX con Word o LibreOffice y se revisa por separado.
"""
import base64
import html
import io
from pathlib import Path
import re

from PIL import Image, ImageOps
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[2]
SOURCE=ROOT/'Documentacion/Informe_Intermedio_Avance_20261020.md'
OUT=ROOT/'Entregables/Informe_Intermedio_Avance_20261020'

def plain(t):
    t=re.sub(r'\[([^]]+)\]\([^)]+\)',r'\1',t)
    return t.replace('**','').replace('`','')

def inline(t):
    t=html.escape(t)
    t=re.sub(r'\*\*(.+?)\*\*',r'<strong>\1</strong>',t)
    t=re.sub(r'`([^`]+)`',r'<code>\1</code>',t)
    return t

def add_text(p,t):
    for j,s in enumerate(re.split(r'(\*\*.*?\*\*|`[^`]+`)',t)):
        run=p.add_run(plain(s))
        if s.startswith('**'):run.bold=True
        if s.startswith('`'):run.font.name='Consolas';run.font.size=Pt(9)

def image_bytes(p):
    im=ImageOps.exif_transpose(Image.open(p)).convert('RGB')
    im.thumbnail((2100,2100))
    b=io.BytesIO();im.save(b,format='JPEG',quality=90);return b.getvalue(),im.size

def table(doc,rows):
    n=len(rows[0]);t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if n==4 and rows[0][0]=='Actividad propuesta':weights=[4.9,3.5,2.2,7.0]
    elif n==4 and rows[0][0]=='Procedencia':weights=[3.8,1.5,1.8,10.5]
    elif n==4 and rows[0][0]=='Destino V2 base':weights=[3.3,2.6,2.4,9.3]
    elif n==4: weights=[2.1,4.2,3.4,6.9]
    elif n==5 and rows[0][0]=='Hito y actividad':weights=[4.7,1.3,2.5,2.5,6.6]
    elif n==5:weights=[4.1,3.6,2.4,2.4,4.1]
    elif n==3 and rows[0][0]=='Tecnología o decisión':weights=[3.5,6.1,8.0]
    elif n==2 and rows[0][0] in ('Concepto','Medida'):weights=[4.4,13.2]
    elif n==8:weights=[4.2,1.75,1.9,1.9,1.9,1.2,2.3,1.45]
    else:weights=[17.6/n]*n
    scale=17.6/sum(weights)
    for j,w in enumerate(weights):t.columns[j].width=Cm(w*scale)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,cell in enumerate(cells):
            cell.width=Cm(weights[j]*scale);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            fill='DFE7EE' if i==0 else ('F6F8FA' if i%2==0 else 'FFFFFF')
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),fill);tcpr.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.05
            txt=plain(row[j]).replace('/', '/\u200b').replace('_','_\u200b')
            r=p.add_run(txt);r.font.size=Pt(8 if n>5 else 9);r.bold=i==0
            if re.fullmatch(r'[\d,./% –]+',plain(row[j])):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def main():
    source=SOURCE.read_text(encoding='utf-8-sig')
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.65);sec.left_margin=Cm(1.7);sec.right_margin=Cm(1.7)
    sec.header_distance=Cm(.6);sec.footer_distance=Cm(.7)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10);normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.12
    for name,size in [('Title',27),('Heading 1',20),('Heading 2',16),('Heading 3',14)]:
        s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s.font.bold=True
        s.paragraph_format.space_before=Pt(7);s.paragraph_format.space_after=Pt(10)
    for border in list(doc.styles.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Walmart Computer Vision  |  20 oct 2026  |  ').font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.title='Informe intermedio de avance Walmart Computer Vision'
    doc.core_properties.author='Vicente Canio, Paulo Guerra y Vicente Vilches'
    css='''body{font:15px/1.5 Arial,sans-serif;color:#151515;background:#e9edf0;margin:0}main{max-width:950px;margin:auto}section{background:white;padding:48px 55px;margin:24px 0}h1{font-size:34px}h2{font-size:25px}h3{font-size:22px}table{width:100%;border-collapse:collapse;margin:20px 0;font-size:13px;overflow-wrap:anywhere}td,th{border:1px solid #d9d9d9;padding:9px;text-align:left}th{background:#dfe7ee}tr:nth-child(even){background:#f6f8fa}figure{margin:16px 0;text-align:center}img{max-width:100%;max-height:400px;object-fit:contain}.original img{max-height:160px}figcaption{font-size:12px;color:#444}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f4f6f8;padding:12px;font-size:12px}@media print{body{background:white}main{max-width:none}section{padding:0;margin:0;break-after:page}section:last-child{break-after:auto}table,figure{break-inside:avoid}@page{size:A4;margin:17mm}}'''
    html_pages=[]
    for pi,page in enumerate(source.split('<!-- page -->')):
        if pi:doc.add_page_break()
        h=[];lines=page.strip().splitlines();i=0;image_index=0
        while i<len(lines):
            line=lines[i].strip()
            if not line:i+=1;continue
            if line.startswith('```'):
                i+=1;code=[]
                while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8)
                r=p.add_run('\n'.join(code));r.font.name='Consolas';r.font.size=Pt(8)
                h.append('<pre>'+html.escape('\n'.join(code))+'</pre>');i+=1;continue
            if line.startswith('|'):
                rows=[]
                while i<len(lines) and lines[i].strip().startswith('|'):
                    raw=lines[i].strip().strip('|');cells=[x.strip() for x in raw.split('|')]
                    if not all(re.fullmatch(r'[:\- ]+',x) for x in cells):rows.append(cells)
                    i+=1
                table(doc,rows)
                h.append('<table>'+''.join('<tr>'+''.join(f'<{"th" if ri==0 else "td"}>'+inline(c) + f'</{"th" if ri==0 else "td"}>' for c in row)+'</tr>' for ri,row in enumerate(rows))+'</table>');continue
            m=re.match(r'!\[([^]]*)\]\(([^)]+)\)',line)
            if m:
                caption,path=m.groups();p=(SOURCE.parent/path).resolve();raw,size=image_bytes(p)
                original='original de entrada' in caption.lower()
                maxw=6.9;maxh=1.35 if original else (4.2 if 'gantt' in caption.lower() else (2.55 if 'arquitectura' in caption.lower() else (2.76 if 'predicción' in caption.lower() else 3.15)))
                w=min(maxw,maxh*size[0]/size[1]);para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_after=Pt(3);para.paragraph_format.keep_with_next=True
                run=para.add_run();sh=run.add_picture(io.BytesIO(raw),width=Cm(w*2.54));sh._inline.docPr.set('descr',caption)
                cp=doc.add_paragraph(caption);cp.alignment=WD_ALIGN_PARAGRAPH.CENTER;cp.paragraph_format.space_after=Pt(7)
                for rr in cp.runs:rr.font.size=Pt(8)
                uri='data:image/jpeg;base64,'+base64.b64encode(raw).decode()
                h.append(f'<figure class="{"original" if original else "comparison"}"><img src="{uri}" alt="{html.escape(caption)}"><figcaption>{html.escape(caption)}</figcaption></figure>');image_index+=1;i+=1;continue
            if line.startswith('#'):
                level=len(line)-len(line.lstrip('#'));text=line[level:].strip()
                p=doc.add_paragraph(plain(text),style='Title' if level==1 else ('Heading 2' if level==2 else 'Heading 3'))
                h.append(f'<h{min(level,3)}>'+inline(text)+f'</h{min(level,3)}>');i+=1;continue
            p=doc.add_paragraph();add_text(p,line);h.append('<p>'+inline(line)+'</p>');i+=1
        html_pages.append('<section>'+'\n'.join(h)+'</section>')
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT.with_suffix('.docx'))
    OUT.with_suffix('.html').write_text('<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>Informe intermedio Walmart Computer Vision</title><style>'+css+'</style><main>'+''.join(html_pages)+'</main></html>',encoding='utf-8')
    print('DOCX y HTML:',OUT,'imágenes incrustadas',sum(1 for x in doc.inline_shapes))

if __name__=='__main__':main()
